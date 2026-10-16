"""Extract: download the Online Retail II dataset and load it into DuckDB staging.

Source: UCI Machine Learning Repository (CC BY 4.0) — the same file that is
mirrored on Kaggle as "Online Retail II". Two sheets cover Dec 2009 – Dec 2011.

The xlsx is streamed to CSV once with openpyxl (read-only mode), then bulk
loaded with DuckDB COPY — orders of magnitude faster than row-by-row inserts.
"""
import csv
import sys
import time
import urllib.request
from pathlib import Path

import duckdb

DB_PATH = "analytics.duckdb"
RAW_DIR = Path("data/raw")
RAW_FILE = RAW_DIR / "online_retail_II.xlsx"
CSV_FILE = RAW_DIR / "online_retail_ii.csv"
SOURCE_URL = "https://archive.ics.uci.edu/ml/machine-learning-databases/00502/online_retail_II.xlsx"

COLUMNS = [
    "invoice", "stock_code", "description", "quantity",
    "invoice_date", "unit_price", "customer_id", "country",
]


def download() -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    if RAW_FILE.exists() and RAW_FILE.stat().st_size > 1_000_000:
        print(f"[extract] using cached {RAW_FILE}", flush=True)
        return
    print(f"[extract] downloading {SOURCE_URL}", flush=True)
    urllib.request.urlretrieve(SOURCE_URL, RAW_FILE)
    print(f"[extract] saved {RAW_FILE} ({RAW_FILE.stat().st_size / 1e6:.1f} MB)", flush=True)


def xlsx_to_csv() -> None:
    if CSV_FILE.exists() and CSV_FILE.stat().st_size > 10_000_000:
        print(f"[extract] using cached {CSV_FILE}", flush=True)
        return

    from openpyxl import load_workbook

    wb = load_workbook(RAW_FILE, read_only=True, data_only=True)
    total = 0
    with open(CSV_FILE, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([*COLUMNS, "source_sheet"])
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            assert "invoice" in str(header[0]).lower(), f"unexpected header: {header}"
            t0 = time.time()
            for r in rows:
                if r[0] is None:
                    continue
                invoice_date = r[4].strftime("%Y-%m-%d %H:%M:%S") if r[4] else ""
                writer.writerow([r[0], r[1], r[2], r[3], invoice_date, r[5], r[6], r[7], sheet])
                total += 1
            print(f"[extract] sheet '{sheet}' done — {total:,} rows "
                  f"({time.time() - t0:.0f}s)", flush=True)
    wb.close()


def load(conn: duckdb.DuckDBPyConnection) -> None:
    conn.execute(
        """
        CREATE OR REPLACE TABLE stg_raw (
            invoice        VARCHAR,
            stock_code     VARCHAR,
            description    VARCHAR,
            quantity       BIGINT,
            invoice_date   TIMESTAMP,
            unit_price     DOUBLE,
            customer_id    BIGINT,
            country        VARCHAR,
            source_sheet   VARCHAR
        )
        """
    )
    conn.execute(
        f"""
        COPY stg_raw FROM '{CSV_FILE.as_posix()}' (
            HEADER, DELIMITER ',', QUOTE '"', NULL ''
        )
        """
    )


def main() -> None:
    download()
    conn = duckdb.connect(DB_PATH)
    try:
        xlsx_to_csv()
        t0 = time.time()
        load(conn)
        n = conn.execute("SELECT COUNT(*) FROM stg_raw").fetchone()[0]
        print(f"[extract] COPY into stg_raw: {time.time() - t0:.0f}s", flush=True)
        print(f"[extract] staging loaded: {n:,} rows -> {DB_PATH}", flush=True)
    finally:
        conn.close()


if __name__ == "__main__":
    sys.exit(main())
